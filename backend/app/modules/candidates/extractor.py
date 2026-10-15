"""
Service for extracting raw plain text from candidate resume files (.pdf and .docx).
Designed to support both manual full-text search and downstream AI parsing pipelines.
"""

from dataclasses import dataclass
from enum import StrEnum
from pathlib import Path

import docx
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.core.exceptions import CVExtractionError


class SupportedFormat(StrEnum):
    PDF = "pdf"
    DOCX = "docx"


@dataclass(frozen=True)
class CVExtractionResult:
    """Normalized output containing extracted text and metadata."""

    text: str
    char_count: int
    word_count: int
    file_format: SupportedFormat


class CVTextExtractorService:
    """Encapsulates parsing logic for PDF and DOCX resume files."""

    def extract_from_path(self, file_path: Path) -> CVExtractionResult:
        """Inspects file extension and delegates to the appropriate parser."""
        if not file_path.exists():
            raise FileNotFoundError(f"File not found at {file_path}")

        ext = file_path.suffix.lower().lstrip(".")

        if ext == SupportedFormat.PDF:
            return self._extract_pdf(file_path)
        elif ext == SupportedFormat.DOCX:
            return self._extract_docx(file_path)
        else:
            raise CVExtractionError(
                f"File extension '{ext}' is not supported for text extraction.",
                code="UNSUPPORTED_FILE_EXTENSION",
            )

    def _extract_pdf(self, file_path: Path) -> CVExtractionResult:
        try:
            reader = PdfReader(file_path)
            if reader.is_encrypted:
                raise CVExtractionError(
                    "The uploaded PDF is password-protected or encrypted.",
                    code="ENCRYPTED_FILE",
                )

            extracted_chunks: list[str] = []
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    extracted_chunks.append(page_text.strip())

            raw_text = "\n\n".join(extracted_chunks).strip()
            if not raw_text:
                raise CVExtractionError(
                    "The uploaded PDF contains no extractable text or is image-only.",
                    code="EMPTY_DOCUMENT",
                )

            return CVExtractionResult(
                text=raw_text,
                char_count=len(raw_text),
                word_count=len(raw_text.split()),
                file_format=SupportedFormat.PDF,
            )
        except (PdfReadError, Exception) as e:
            if isinstance(e, CVExtractionError):
                raise
            raise CVExtractionError(
                f"Failed to parse PDF file: {e!s}",
                code="CORRUPTED_FILE",
            )

    def _extract_docx(self, file_path: Path) -> CVExtractionResult:
        try:
            doc = docx.Document(file_path)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

            # Extract text from tables (e.g. skills or experience grids)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        if cell_text and cell_text not in paragraphs:
                            paragraphs.append(cell_text)

            raw_text = "\n".join(paragraphs).strip()
            if not raw_text:
                raise CVExtractionError(
                    "The uploaded DOCX document contains no extractable text.",
                    code="EMPTY_DOCUMENT",
                )

            return CVExtractionResult(
                text=raw_text,
                char_count=len(raw_text),
                word_count=len(raw_text.split()),
                file_format=SupportedFormat.DOCX,
            )
        except Exception as e:
            if isinstance(e, CVExtractionError):
                raise
            raise CVExtractionError(
                f"Failed to parse DOCX file: {e!s}",
                code="CORRUPTED_FILE",
            )
