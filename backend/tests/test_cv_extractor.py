"""
Unit tests for CVTextExtractorService, covering valid DOCX extraction
and error handling for empty documents, unsupported extensions, and corrupted files.
"""

from pathlib import Path

import docx
import pytest

from app.core.exceptions import CVExtractionError
from app.modules.candidates.extractor import (
    CVExtractionResult,
    CVTextExtractorService,
    SupportedFormat,
)


@pytest.fixture
def extractor() -> CVTextExtractorService:
    return CVTextExtractorService()


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    """Generates a valid, readable DOCX resume file."""
    docx_path = tmp_path / "sample_resume.docx"
    doc = docx.Document()
    doc.add_heading("Muhammad Riaz", level=1)
    doc.add_paragraph("Senior Mobile and Full-Stack Software Engineer.")

    # Add a table to test table text extraction
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "Flutter, Python, FastAPI, PostgreSQL"

    doc.save(docx_path)
    return docx_path


@pytest.fixture
def empty_docx(tmp_path: Path) -> Path:
    """Generates an empty DOCX file containing no text."""
    docx_path = tmp_path / "empty_resume.docx"
    doc = docx.Document()
    doc.save(docx_path)
    return docx_path


@pytest.fixture
def corrupted_file(tmp_path: Path) -> Path:
    """Generates a malformed non-binary file with a .pdf extension."""
    file_path = tmp_path / "corrupted.pdf"
    file_path.write_text("NOT A REAL PDF CONTENT HEADER")
    return file_path


# --- Test Cases ---


def test_extract_docx_success(
    extractor: CVTextExtractorService, sample_docx: Path
) -> None:
    result = extractor.extract_from_path(sample_docx)

    assert isinstance(result, CVExtractionResult)
    assert result.file_format == SupportedFormat.DOCX
    assert "Muhammad Riaz" in result.text
    assert "Senior Mobile and Full-Stack Software Engineer" in result.text
    assert "Flutter, Python, FastAPI, PostgreSQL" in result.text
    assert result.word_count > 0
    assert result.char_count > 0


def test_extract_empty_docx_raises_error(
    extractor: CVTextExtractorService, empty_docx: Path
) -> None:
    with pytest.raises(CVExtractionError) as exc_info:
        extractor.extract_from_path(empty_docx)

    assert exc_info.value.code == "EMPTY_DOCUMENT"
    assert exc_info.value.status_code == 400


def test_extract_unsupported_extension_raises_error(
    extractor: CVTextExtractorService, tmp_path: Path
) -> None:
    txt_file = tmp_path / "resume.txt"
    txt_file.write_text("Plain text content")

    with pytest.raises(CVExtractionError) as exc_info:
        extractor.extract_from_path(txt_file)

    assert exc_info.value.code == "UNSUPPORTED_FILE_EXTENSION"
    assert exc_info.value.status_code == 400


def test_extract_corrupted_pdf_raises_error(
    extractor: CVTextExtractorService, corrupted_file: Path
) -> None:
    with pytest.raises(CVExtractionError) as exc_info:
        extractor.extract_from_path(corrupted_file)

    assert exc_info.value.code == "CORRUPTED_FILE"
    assert exc_info.value.status_code == 400


def test_extract_missing_file_raises_file_not_found(
    extractor: CVTextExtractorService, tmp_path: Path
) -> None:
    missing_path = tmp_path / "non_existent.pdf"

    with pytest.raises(FileNotFoundError):
        extractor.extract_from_path(missing_path)
